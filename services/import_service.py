# services/import_service.py
# ============================================================
#  Import produits en masse — Excel, CSV, PDF
# ============================================================

import io
import re
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

import pandas as pd

COLONNES_REQUISES  = ["code_produit", "designation", "categorie"]
COLONNES_OPTIONNELLES = [
    "unite", "prix_achat", "prix_vente",
    "stock", "stock_minimum", "emplacement", "statut",
]
TOUTES_COLONNES = COLONNES_REQUISES + COLONNES_OPTIONNELLES

CATEGORIES_VALIDES = [
    "Ciment", "Fer", "Peinture", "Electricite", "Électricité",
    "Plomberie", "Outillage", "Visserie", "Serrurerie", "Autre",
]
UNITES_VALIDES = [
    "Piece", "Pièce", "Sac", "Metre", "Mètre",
    "Rouleau", "Bidon", "Barre", "Boite", "Boîte", "Kg",
]
STATUTS_VALIDES = ["Actif", "Inactif"]


# ── Normalisation ─────────────────────────────────────────────────────────

def _normalise_colonne(nom: str) -> str:
    """Normalise un nom de colonne : minuscules, espaces → _, accents retirés."""
    nom = nom.strip().lower()
    replacements = {
        "é": "e", "è": "e", "ê": "e", "ë": "e",
        "à": "a", "â": "a", "ù": "u", "û": "u",
        "î": "i", "ï": "i", "ô": "o", "ç": "c",
    }
    for src, dst in replacements.items():
        nom = nom.replace(src, dst)
    nom = re.sub(r"[\s\-/]+", "_", nom)
    # Aliases courants
    aliases = {
        "code":           "code_produit",
        "ref":            "code_produit",
        "reference":      "code_produit",
        "nom":            "designation",
        "libelle":        "designation",
        "produit":        "designation",
        "cat":            "categorie",
        "unite_mesure":   "unite",
        "prix_achat_ht":  "prix_achat",
        "prix_vente_ttc": "prix_vente",
        "quantite":       "stock",
        "qte":            "stock",
        "stock_actuel":   "stock",
        "seuil":          "stock_minimum",
        "min":            "stock_minimum",
        "emplacement":    "emplacement",
        "location":       "emplacement",
        "zone":           "emplacement",
        "statut":         "statut",
        "etat":           "statut",
    }
    return aliases.get(nom, nom)


def _to_decimal(val: Any, default: Decimal = Decimal(0)) -> Decimal:
    try:
        if pd.isna(val):
            return default
        return Decimal(str(val).replace(",", ".").strip())
    except (InvalidOperation, ValueError):
        return default


def _to_int(val: Any, default: int = 0) -> int:
    try:
        if pd.isna(val):
            return default
        return int(float(str(val).replace(",", ".").strip()))
    except (ValueError, TypeError):
        return default


# ── Lecture des fichiers ──────────────────────────────────────────────────

def lire_excel(file_bytes: bytes) -> pd.DataFrame:
    return pd.read_excel(io.BytesIO(file_bytes), dtype=str)


def lire_csv(file_bytes: bytes) -> pd.DataFrame:
    """Détecte automatiquement le séparateur (virgule, point-virgule, tabulation)."""
    text = file_bytes.decode("utf-8-sig", errors="replace")
    # Compte les séparateurs dans la première ligne
    first_line = text.split("\n")[0]
    sep = ";"
    if first_line.count(",") > first_line.count(";"):
        sep = ","
    elif first_line.count("\t") > first_line.count(";"):
        sep = "\t"
    return pd.read_csv(io.StringIO(text), sep=sep, dtype=str)


def lire_pdf(file_bytes: bytes) -> pd.DataFrame:
    """
    Extrait les tableaux d'un PDF catalogue.
    Fonctionne si le PDF contient des tableaux structurés.
    """
    try:
        import pdfplumber
        rows = []
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            for page in pdf.pages:
                tables = page.extract_tables()
                for table in tables:
                    if not table:
                        continue
                    # Première ligne = en-têtes
                    headers = [str(h).strip() if h else "" for h in table[0]]
                    for row in table[1:]:
                        if any(cell and str(cell).strip() for cell in row):
                            rows.append(dict(zip(headers, [
                                str(c).strip() if c else ""
                                for c in row
                            ])))
        if rows:
            return pd.DataFrame(rows)
        return pd.DataFrame()
    except ImportError:
        raise ImportError("pdfplumber requis pour l'import PDF. "
                         "Installez-le : pip install pdfplumber")
    except Exception as e:
        raise ValueError(f"Impossible de lire le PDF : {e}")


# ── Normalisation du DataFrame ────────────────────────────────────────────

def normaliser_dataframe(df: pd.DataFrame) -> pd.DataFrame:
    """Renomme les colonnes et applique les valeurs par défaut."""
    # Renomme les colonnes
    df.columns = [_normalise_colonne(c) for c in df.columns]

    # Garde seulement les colonnes connues
    cols_presentes = [c for c in TOUTES_COLONNES if c in df.columns]
    df = df[cols_presentes].copy()

    # Valeurs par défaut
    defaults = {
        "unite":         "Pièce",
        "prix_achat":    "0",
        "prix_vente":    "0",
        "stock":         "0",
        "stock_minimum": "5",
        "emplacement":   "",
        "statut":        "Actif",
    }
    for col, default in defaults.items():
        if col not in df.columns:
            df[col] = default

    # Nettoyage
    for col in ["code_produit", "designation", "categorie"]:
        if col in df.columns:
            df[col] = df[col].astype(str).str.strip()

    # Supprime les lignes vides
    df = df[df["code_produit"].notna() & (df["code_produit"] != "") &
            (df["code_produit"] != "nan")]

    # Majuscules sur le code
    df["code_produit"] = df["code_produit"].str.upper()

    return df.reset_index(drop=True)


# ── Validation ────────────────────────────────────────────────────────────

def valider_dataframe(df: pd.DataFrame, session) -> dict:
    """
    Valide chaque ligne et retourne :
    {
        "valides":   liste des dicts prêts à insérer,
        "erreurs":   liste des {ligne, code, erreur},
        "doublons_fichier": codes dupliqués dans le fichier,
        "doublons_db":      codes déjà en base,
        "stats": {...}
    }
    """
    from database.models import Produit
    from sqlalchemy import select

    # Codes déjà en base
    codes_db = set(
        row[0] for row in session.execute(
            select(Produit.code_produit)
        ).all()
    )

    valides          = []
    erreurs          = []
    codes_vus        = set()
    doublons_fichier = set()
    doublons_db      = set()

    for idx, row in df.iterrows():
        ligne_num = idx + 2  # +2 car ligne 1 = en-têtes
        code      = str(row.get("code_produit", "")).strip().upper()
        errs      = []

        # Code obligatoire
        if not code or code == "NAN":
            errs.append("Code produit manquant")

        # Désignation obligatoire
        desig = str(row.get("designation", "")).strip()
        if not desig or desig == "nan":
            errs.append("Désignation manquante")

        # Catégorie obligatoire
        cat = str(row.get("categorie", "")).strip()
        if not cat or cat == "nan":
            errs.append("Catégorie manquante")

        # Doublon dans le fichier
        if code in codes_vus:
            doublons_fichier.add(code)
            errs.append(f"Code dupliqué dans le fichier")
        codes_vus.add(code)

        # Doublon en base
        if code in codes_db:
            doublons_db.add(code)
            # Pas une erreur bloquante — on signale seulement

        # Prix valides
        pa = _to_decimal(row.get("prix_achat", 0))
        pv = _to_decimal(row.get("prix_vente", 0))
        if pa < 0:
            errs.append("Prix achat négatif")
        if pv < 0:
            errs.append("Prix vente négatif")

        # Stock valide
        stock = _to_int(row.get("stock", 0))
        if stock < 0:
            errs.append("Stock négatif")

        if errs:
            erreurs.append({
                "Ligne": ligne_num,
                "Code":  code,
                "Erreur": " | ".join(errs),
            })
        else:
            valides.append({
                "code_produit":  code,
                "designation":   desig,
                "categorie":     cat,
                "unite":         str(row.get("unite", "Pièce")).strip() or "Pièce",
                "prix_achat":    pa,
                "prix_vente":    pv,
                "stock":         stock,
                "stock_minimum": _to_int(row.get("stock_minimum", 5)),
                "marge":         pv - pa,
                "emplacement":   str(row.get("emplacement", "")).strip() or None,
                "statut":        str(row.get("statut", "Actif")).strip() or "Actif",
            })

    return {
        "valides":          valides,
        "erreurs":          erreurs,
        "doublons_fichier": doublons_fichier,
        "doublons_db":      doublons_db,
        "stats": {
            "total":     len(df),
            "valides":   len(valides),
            "erreurs":   len(erreurs),
            "doublons_db": len(doublons_db),
        },
    }


# ── Import en base ────────────────────────────────────────────────────────

def importer_produits(session, valides: list,
                      doublons_db: set,
                      mode_doublon: str = "ignorer") -> dict:
    """
    Insère les produits valides.
    mode_doublon : 'ignorer' | 'mettre_a_jour'
    Retourne {inseres, mis_a_jour, ignores}
    """
    from database.models import Produit
    from services.logging_service import log_action

    inseres     = 0
    mis_a_jour  = 0
    ignores     = 0

    for item in valides:
        code = item["code_produit"]

        if code in doublons_db:
            if mode_doublon == "mettre_a_jour":
                p = session.get(Produit, code)
                if p:
                    p.designation   = item["designation"]
                    p.categorie     = item["categorie"]
                    p.unite         = item["unite"]
                    p.prix_achat    = item["prix_achat"]
                    p.prix_vente    = item["prix_vente"]
                    p.stock         = item["stock"]
                    p.stock_minimum = item["stock_minimum"]
                    p.marge         = item["marge"]
                    p.emplacement   = item["emplacement"]
                    p.statut        = item["statut"]
                    mis_a_jour += 1
            else:
                ignores += 1
        else:
            session.add(Produit(**item))
            inseres += 1

    session.commit()
    log_action(session, None,
               f"Import produits : {inseres} insérés, "
               f"{mis_a_jour} mis à jour, {ignores} ignorés")

    return {"inseres": inseres, "mis_a_jour": mis_a_jour, "ignores": ignores}


# ── Génération template Excel ─────────────────────────────────────────────

def generer_template_excel() -> bytes:
    """Génère un fichier Excel template avec exemples et commentaires."""
    import openpyxl
    from openpyxl.styles import (Font, PatternFill, Alignment,
                                  Border, Side, numbers)
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Produits"

    # En-têtes
    headers = [
        "code_produit", "designation", "categorie", "unite",
        "prix_achat", "prix_vente", "stock", "stock_minimum",
        "emplacement", "statut",
    ]
    descriptions = [
        "Code unique *", "Nom du produit *", "Catégorie *", "Unité de vente",
        "Prix achat", "Prix vente", "Stock initial", "Stock minimum",
        "Emplacement", "Actif/Inactif",
    ]

    # Style en-têtes
    header_fill = PatternFill("solid", fgColor="1E3A5F")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    desc_fill   = PatternFill("solid", fgColor="EFF6FF")
    desc_font   = Font(color="1E3A5F", italic=True, size=9)

    for col_idx, (h, d) in enumerate(zip(headers, descriptions), 1):
        # Ligne 1 : noms de colonnes
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

        # Ligne 2 : descriptions
        cell2 = ws.cell(row=2, column=col_idx, value=d)
        cell2.fill = desc_fill
        cell2.font = desc_font
        cell2.alignment = Alignment(horizontal="center")

    # Exemples de données
    exemples = [
        ["CIM-001", "Ciment CEM II 50kg",       "Ciment",      "Sac",     4500, 5500, 100, 20, "Zone A-1", "Actif"],
        ["FER-001", "Fer HA 12mm 12m",           "Fer",         "Barre",   8200, 11000, 50,  10, "Zone B-1", "Actif"],
        ["PEI-001", "Peinture Blanche 10L",      "Peinture",    "Bidon",  18000, 24500, 30,   5, "Zone C-1", "Actif"],
        ["ELE-001", "Cable 2.5mm 100m",          "Electricite", "Rouleau",24000, 32000, 20,   5, "Zone D-1", "Actif"],
        ["PLO-001", "Robinet ball valve 1/2",    "Plomberie",   "Piece",   1800,  2800, 50,  10, "Zone E-1", "Actif"],
        ["VIS-001", "Cheville Fischer 8mm x100", "Visserie",    "Boite",    900,  1500,200,  50, "Zone F-1", "Actif"],
    ]

    row_fill_1 = PatternFill("solid", fgColor="FFFFFF")
    row_fill_2 = PatternFill("solid", fgColor="F8FAFC")

    for row_idx, exemple in enumerate(exemples, 3):
        fill = row_fill_1 if row_idx % 2 == 0 else row_fill_2
        for col_idx, val in enumerate(exemple, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.fill = fill
            cell.alignment = Alignment(vertical="center")

    # Largeurs colonnes
    widths = [14, 28, 14, 10, 12, 12, 8, 14, 12, 10]
    for col_idx, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = w

    # Hauteur lignes
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 18
    for i in range(3, 3 + len(exemples)):
        ws.row_dimensions[i].height = 18

    # Feuille de référence
    ws2 = wb.create_sheet("Reference")
    ws2["A1"] = "CATEGORIES VALIDES"
    ws2["A1"].font = Font(bold=True)
    cats = ["Ciment","Fer","Peinture","Electricite","Plomberie",
            "Outillage","Visserie","Serrurerie","Autre"]
    for i, c in enumerate(cats, 2):
        ws2.cell(row=i, column=1, value=c)

    ws2["C1"] = "UNITES VALIDES"
    ws2["C1"].font = Font(bold=True)
    unites = ["Piece","Sac","Metre","Rouleau","Bidon","Barre","Boite","Kg"]
    for i, u in enumerate(unites, 2):
        ws2.cell(row=i, column=3, value=u)

    ws2["E1"] = "STATUTS VALIDES"
    ws2["E1"].font = Font(bold=True)
    ws2["E2"] = "Actif"
    ws2["E3"] = "Inactif"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()